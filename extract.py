"""Load files, run ContextGem against Gemini Flash, serialize results."""

from __future__ import annotations

import io
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from contextgem import Document, DocumentLLM, ExtractionPipeline


FLASH_MODELS = {
    "Gemini 3.8 Flash (recommended)": "gemini/gemini-3.8-flash",
    "Gemini Flash latest alias": "gemini/gemini-flash-latest",
    "Gemini 2.5 Flash": "gemini/gemini-2.5-flash",
    "Gemini 3.6 Flash": "gemini/gemini-3.6-flash",
}


def text_from_upload(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json", ".html", ".xml"}:
        return data.decode("utf-8", errors="replace")

    if suffix == ".docx":
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            try:
                from contextgem import DocxConverter

                return DocxConverter().convert_to_text_format(
                    tmp.name, output_format="markdown"
                )
            except Exception:
                try:
                    import docx

                    parsed = docx.Document(tmp.name)
                    return "\n".join(
                        p.text for p in parsed.paragraphs if p.text.strip()
                    )
                except Exception as exc:
                    raise RuntimeError(f"Could not read DOCX: {exc}") from exc

    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            pages = []
            for i, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                pages.append(f"--- Page {i} ---\n{text}")
            return "\n\n".join(pages).strip()
        except Exception as exc:
            raise RuntimeError(f"Could not read PDF: {exc}") from exc

    raise RuntimeError(
        f"Unsupported file type '{suffix}'. Use .txt, .md, .docx, or .pdf."
    )


def make_llm(api_key: str, model: str) -> DocumentLLM:
    if not api_key or not api_key.strip():
        raise ValueError("Paste a Google AI Studio API key first.")
    if "/" not in model:
        model = f"gemini/{model}"
    return DocumentLLM(
        model=model,
        api_key=api_key.strip(),
        role="extractor_text",
        temperature=0.2,
        timeout=180,
        max_retries_invalid_data=2,
    )


def run_extraction(text: str, pipeline: ExtractionPipeline, llm: DocumentLLM) -> Document:
    cleaned = (text or "").strip()
    if len(cleaned) < 20:
        raise ValueError("Document is empty or too short.")
    doc = Document(raw_text=cleaned)
    doc.assign_pipeline(pipeline)
    return llm.extract_all(doc)


def _serialize_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (list, tuple)):
        return [_serialize_value(v) for v in value]
    if isinstance(value, dict):
        return {k: _serialize_value(v) for k, v in value.items()}
    return value


def _item_to_dict(item: Any) -> dict[str, Any]:
    payload: dict[str, Any] = {"value": _serialize_value(getattr(item, "value", None))}
    justification = getattr(item, "justification", None)
    if justification:
        payload["justification"] = justification

    refs: list[str] = []
    for sent in getattr(item, "reference_sentences", None) or []:
        raw = getattr(sent, "raw_text", None)
        if raw:
            refs.append(raw)
    if not refs:
        for para in getattr(item, "reference_paragraphs", None) or []:
            raw = getattr(para, "raw_text", None)
            if raw:
                refs.append(raw)
    if refs:
        payload["sources"] = refs
    return payload


def document_to_results(doc: Document) -> dict[str, Any]:
    concepts = []
    for concept in doc.concepts or []:
        concepts.append(
            {
                "name": concept.name,
                "description": getattr(concept, "description", ""),
                "items": [_item_to_dict(item) for item in (concept.extracted_items or [])],
            }
        )

    aspects = []
    for aspect in doc.aspects or []:
        aspect_payload: dict[str, Any] = {
            "name": aspect.name,
            "description": getattr(aspect, "description", ""),
            "items": [_item_to_dict(item) for item in (aspect.extracted_items or [])],
            "concepts": [],
        }
        for concept in getattr(aspect, "concepts", None) or []:
            aspect_payload["concepts"].append(
                {
                    "name": concept.name,
                    "description": getattr(concept, "description", ""),
                    "items": [_item_to_dict(item) for item in (concept.extracted_items or [])],
                }
            )
        aspects.append(aspect_payload)

    return {"concepts": concepts, "aspects": aspects}
