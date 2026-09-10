import io
import json
import tempfile
from datetime import date, datetime
from pathlib import Path

import streamlit as st
from contextgem import (
    Aspect,
    BooleanConcept,
    DateConcept,
    Document,
    DocumentLLM,
    ExtractionPipeline,
    NumericalConcept,
    StringConcept,
)

st.set_page_config(page_title="ContextGem · Gemini Flash", page_icon="💎", layout="wide")
st.title("💎 ContextGem")
st.caption("Structured extraction with Gemini Flash. Key stays in this session.")


def string_field(name, description, singular=False):
    return StringConcept(
        name=name,
        description=description,
        add_references=True,
        reference_depth="sentences",
        add_justifications=True,
        justification_depth="brief",
        singular_occurrence=singular,
    )


PRESETS = {
    "Contract / agreement": ExtractionPipeline(
        aspects=[
            Aspect(name="Parties", description="Parties and their roles.", reference_depth="sentences"),
            Aspect(name="Commercial terms", description="Fees, payment, commercial obligations.", reference_depth="sentences"),
            Aspect(name="Term and termination", description="Duration, renewal, termination.", reference_depth="sentences"),
        ],
        concepts=[
            string_field("Document type", "Type of agreement.", True),
            string_field("Party A", "First party name and role.", True),
            string_field("Party B", "Second party name and role.", True),
            DateConcept(name="Effective date", description="Start or effective date.", add_references=True, singular_occurrence=True),
            string_field("Governing law", "Governing jurisdiction.", True),
            string_field("Termination notice", "Notice period to terminate.", True),
            string_field("Payment terms", "When and how payment is due."),
            BooleanConcept(name="Auto-renewal", description="Renews automatically unless cancelled.", add_justifications=True, singular_occurrence=True),
        ],
    ),
    "Invoice / bill": ExtractionPipeline(
        concepts=[
            string_field("Supplier", "Vendor name.", True),
            string_field("Customer", "Bill-to name.", True),
            string_field("Invoice number", "Invoice id.", True),
            DateConcept(name="Invoice date", description="Date on the invoice.", add_references=True, singular_occurrence=True),
            DateConcept(name="Due date", description="Payment due date.", add_references=True, singular_occurrence=True),
            NumericalConcept(name="Total", description="Total amount due.", numeric_type="float", singular_occurrence=True),
            string_field("Currency", "Currency code or symbol.", True),
            string_field("Line items", "Each product or service line."),
        ]
    ),
    "General document": ExtractionPipeline(
        concepts=[
            string_field("Title", "Document title.", True),
            string_field("Key people or organizations", "Named people or organizations."),
            DateConcept(name="Important dates", description="Dates that matter.", add_references=True),
            string_field("Key facts", "Important facts or conclusions."),
            string_field("Action items", "Tasks or next steps."),
        ]
    ),
}


def read_upload(name, data):
    suffix = Path(name).suffix.lower()
    if suffix in {".txt", ".md", ".csv"}:
        return data.decode("utf-8", errors="replace")
    if suffix == ".docx":
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            try:
                from contextgem import DocxConverter
                return DocxConverter().convert_to_text_format(tmp.name, output_format="markdown")
            except Exception:
                import docx
                doc = docx.Document(tmp.name)
                return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    raise RuntimeError(f"Unsupported type: {suffix}")


def item_dict(item):
    payload = {"value": item.value.isoformat() if isinstance(item.value, (date, datetime)) else item.value}
    if getattr(item, "justification", None):
        payload["justification"] = item.justification
    sources = []
    for sent in getattr(item, "reference_sentences", None) or []:
        if getattr(sent, "raw_text", None):
            sources.append(sent.raw_text)
    if sources:
        payload["sources"] = sources
    return payload


with st.sidebar:
    api_key = st.text_input("Google AI Studio API key", type="password")
    model = st.selectbox(
        "Flash model",
        ["gemini/gemini-3.8-flash", "gemini/gemini-2.5-flash", "gemini/gemini-flash-latest"],
    )

col1, col2 = st.columns(2)
with col1:
    source = st.radio("Input", ["Paste text", "Upload file"], horizontal=True)
    text = ""
    if source == "Paste text":
        text = st.text_area("Document", height=240)
    else:
        uploaded = st.file_uploader("TXT / MD / DOCX / PDF", type=["txt", "md", "docx", "pdf"])
        if uploaded:
            text = read_upload(uploaded.name, uploaded.getvalue())
            st.success(f"Loaded {len(text):,} characters")
    preset = st.selectbox("Preset", list(PRESETS))
    run = st.button("Extract with Gemini Flash", type="primary")

with col2:
    if run:
        if not api_key:
            st.error("Paste an API key from https://aistudio.google.com/apikey")
        elif len((text or "").strip()) < 20:
            st.error("Document is empty.")
        else:
            with st.spinner("Gemini Flash is reading the document…"):
                llm = DocumentLLM(model=model, api_key=api_key.strip(), temperature=0.2, timeout=180)
                doc = Document(raw_text=text.strip())
                doc.assign_pipeline(PRESETS[preset])
                doc = llm.extract_all(doc)
            rows = []
            for concept in doc.concepts or []:
                items = concept.extracted_items or []
                if not items:
                    rows.append({"Field": concept.name, "Value": "—", "Why": ""})
                for item in items:
                    data = item_dict(item)
                    rows.append(
                        {
                            "Field": concept.name,
                            "Value": data.get("value"),
                            "Why": data.get("justification", ""),
                        }
                    )
            if rows:
                st.dataframe(rows, use_container_width=True, hide_index=True)
            for aspect in doc.aspects or []:
                st.markdown(f"**{aspect.name}**")
                for item in aspect.extracted_items or []:
                    st.write(item.value)
            st.download_button(
                "Download JSON",
                json.dumps(rows, indent=2, default=str),
                "extract.json",
                "application/json",
            )
    else:
        st.info("Add a document and click Extract.")
